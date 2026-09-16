import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from knowledge_base.docx_normalizer import normalize_docx


class DocxNormalizerTests(unittest.TestCase):
    def test_splits_manual_lines_inside_paragraph(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "lines.docx"

            document = Document()
            document.add_paragraph(
                "Первая строка\nВторая строка\nТретья строка"
            )
            document.save(file_path)

            lines = normalize_docx(file_path)

            self.assertEqual(len(lines), 3)

            self.assertEqual(
                [line.text for line in lines],
                [
                    "Первая строка",
                    "Вторая строка",
                    "Третья строка",
                ],
            )

            self.assertEqual(
                {line.paragraph_index for line in lines},
                {0},
            )

    def test_preserves_formatting_signals(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "heading.docx"

            document = Document()

            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = paragraph.add_run("Федеральные законы")
            run.bold = True
            run.font.size = Pt(14)

            document.save(file_path)

            lines = normalize_docx(file_path)

            self.assertEqual(len(lines), 1)

            line = lines[0]

            self.assertEqual(
                line.text,
                "Федеральные законы",
            )

            self.assertEqual(
                line.bold_percent,
                100,
            )

            self.assertEqual(
                line.font_sizes,
                (14.0,),
            )

            self.assertEqual(
                line.alignment,
                "CENTER",
            )

    def test_manual_lines_keep_individual_formatting(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "mixed.docx"

            document = Document()

            paragraph = document.add_paragraph()

            first = paragraph.add_run("Жирная строка")
            first.bold = True
            first.font.size = Pt(14)

            paragraph.add_run("\n")

            second = paragraph.add_run("Обычная строка")
            second.bold = False
            second.font.size = Pt(10)

            document.save(file_path)

            lines = normalize_docx(file_path)

            self.assertEqual(len(lines), 2)

            self.assertEqual(lines[0].bold_percent, 100)
            self.assertEqual(lines[0].font_sizes, (14.0,))

            self.assertEqual(lines[1].bold_percent, 0)
            self.assertEqual(lines[1].font_sizes, (10.0,))

    def test_blank_lines_are_removed(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "blank.docx"

            document = Document()
            document.add_paragraph(
                "Первая\n\n   \nВторая"
            )
            document.save(file_path)

            lines = normalize_docx(file_path)

            self.assertEqual(
                [line.text for line in lines],
                ["Первая", "Вторая"],
            )


if __name__ == "__main__":
    unittest.main()
