"""Real SQLite and real DOCX/PDF ingestion, using temporary synthetic files only."""
from __future__ import annotations

import ast
from dataclasses import replace
import hashlib
import inspect
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document

from knowledge_base.runtime.ingestion import IngestionError, IngestionService
from knowledge_base.runtime.models import (
    AccessContext, ChunkIdentity, DocumentIdentity, DocumentVersionIdentity,
    DocumentVersionStatus, ProcessingSnapshotIdentity, ProcessingState,
    RetrievalFilters, RuntimeDocument, RuntimeDocumentVersion,
)
from knowledge_base.runtime.protocols import ChunkRepository, DocumentRepository
from knowledge_base.runtime.storage import SQLiteRepository, StorageError
from test_pdf_parser import create_simple_pdf

ROOT = Path(__file__).resolve().parents[1]


class RuntimeIngestionTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.path = Path(self.folder.name)
        self.database = self.path / 'registry.sqlite'
        self.repository = SQLiteRepository(self.database)
        self.addCleanup(self.repository.close)
        self.service = IngestionService(self.repository, max_chunk_chars=100, overlap_pieces=0)

    def docx(self, name='input.docx', texts=('Supply requests go to the fictional company desk.',)):
        path = self.path / name
        document = Document()
        for text in texts:
            document.add_paragraph(text)
        document.save(path)
        return path

    def ingest(self, file=None, **kwargs):
        return self.service.ingest_new_document(file or self.docx(), title='Independent document',
                                               organization_id='org/any', version_label='revision/first', **kwargs)

    def publish(self, result, previous=None):
        self.service.activate_version(result.version.identity, expected_current_version_id=previous)

    def table_count(self, table):
        # Fixed test table names, never caller-supplied SQL.
        return self.repository.connection.execute('SELECT COUNT(*) FROM ' + table).fetchone()[0]

    def test_adapter_satisfies_both_existing_repository_protocols(self):
        self.assertIsInstance(self.repository, DocumentRepository)
        self.assertIsInstance(self.repository, ChunkRepository)
        for contract in (DocumentRepository, ChunkRepository):
            for name, member in vars(contract).items():
                if inspect.isfunction(member) and not name.startswith('_'):
                    self.assertEqual(list(inspect.signature(member).parameters),
                                     list(inspect.signature(getattr(SQLiteRepository, name)).parameters))

    def test_new_document_source_version_and_chunks_persist(self):
        file = self.docx(texts=('Unique delivery rules.', 'A second source paragraph.'))
        result = self.ingest(file, metadata={'nested': {'tags': ['independent']}})
        self.assertFalse(result.duplicate)
        self.assertEqual(result.version.status, DocumentVersionStatus.DRAFT)
        self.assertEqual(result.version.processing_state, ProcessingState.CHUNKED)
        self.assertFalse(result.version.approved)
        self.assertEqual(self.repository.get_document(result.document.identity), result.document)
        self.assertEqual(self.repository.get_version(result.version.identity), result.version)
        self.assertEqual(self.repository.get_source(result.version.identity), file.read_bytes())
        info = self.repository.get_source_info(result.version.identity)
        self.assertEqual(info.sha256, hashlib.sha256(file.read_bytes()).hexdigest())
        self.assertIsNotNone(info.created_at.utcoffset())
        chunks = self.repository.list_snapshot(result.snapshot)
        self.assertEqual(len(chunks), result.chunk_count)
        self.assertIn('Unique delivery rules.', chunks[0].text)
        self.assertEqual(chunks[0].coordinates.paragraph_indices, (0, 1))
        self.assertEqual(self.repository.get_snapshot_metadata(result.snapshot)['max_chunk_chars'], 100)

    def test_two_different_corpora_use_same_real_pipeline_and_disjoint_ids(self):
        company = self.docx('company.docx', ('Travel approval is handled by the Cedar desk.',))
        laboratory = self.path / 'lab.pdf'
        create_simple_pdf(laboratory, ['Glass vessels carry a blue label.',
                                      'Polymer vessels carry a silver label.',
                                      'Record the label before storage.'])
        a = self.ingest(company, document_id='company/arbitrary Ω')
        b = self.service.ingest_new_document(laboratory, title='Fictional lab instructions',
            organization_id='org/lab', version_label='lab/next', document_id='lab/id#independent')
        ca, cb = self.repository.list_snapshot(a.snapshot), self.repository.list_snapshot(b.snapshot)
        self.assertNotEqual(len(ca), len(cb))
        self.assertEqual(ca[0].local_chunk_index, cb[0].local_chunk_index)
        self.assertEqual(ca[0].local_chunk_index, 0)
        self.assertFalse({c.identity.chunk_id for c in ca} & {c.identity.chunk_id for c in cb})
        self.assertNotEqual(a.document.identity, b.document.identity)
        self.assertNotEqual(a.version.identity, b.version.identity)
        self.assertEqual(cb[0].coordinates.page_numbers, (1, 2))
        for chunk in (*ca, *cb):
            self.assertEqual(self.repository.get_chunk(chunk.identity), chunk)

    def test_byte_identical_upload_is_reported_without_new_records(self):
        file = self.docx()
        first = self.ingest(file)
        renamed = self.path / 'renamed.docx'
        renamed.write_bytes(file.read_bytes())
        again = self.ingest(renamed)
        self.assertTrue(again.duplicate)
        self.assertEqual(again.version.identity, first.version.identity)
        self.assertEqual(again.snapshot, first.snapshot)
        self.assertEqual(self.table_count('documents'), 1)
        self.assertEqual(self.table_count('versions'), 1)
        self.assertEqual(self.table_count('snapshots'), 1)

    def test_duplicate_existing_version_does_not_create_new_label_or_snapshot(self):
        file = self.docx()
        first = self.ingest(file)
        self.publish(first)
        duplicate = self.service.ingest_new_version(file, document=first.document.identity, version_label='different-label')
        self.assertTrue(duplicate.duplicate)
        self.assertEqual(duplicate.version.identity, first.version.identity)
        self.assertEqual(duplicate.version.status, DocumentVersionStatus.ACTIVE)
        self.assertEqual(duplicate.snapshot, first.snapshot)
        self.assertEqual(self.table_count('versions'), 1)

    def test_explicit_identities_can_intentionally_share_source_hash(self):
        file = self.docx()
        a = self.ingest(file, document_id='explicit/A')
        b = self.ingest(file, document_id='explicit/B')
        self.assertFalse(b.duplicate)
        self.assertNotEqual(a.document.identity, b.document.identity)
        self.assertEqual(self.repository.get_source_info(a.version.identity).sha256,
                         self.repository.get_source_info(b.version.identity).sha256)
        with self.assertRaises(IngestionError) as error:
            self.ingest(file)
        self.assertEqual(error.exception.reason_code, 'ambiguous_duplicate_requires_identity')
        self.assertEqual(self.table_count('documents'), 2)

    def test_hash_lookup_is_scoped_to_organization(self):
        file = self.docx()
        a = self.ingest(file)
        b = self.service.ingest_new_document(file, title='Separate organization', organization_id='org/other', version_label='v')
        self.assertFalse(b.duplicate)
        self.assertNotEqual(a.document.identity.document_id, b.document.identity.document_id)
        self.assertEqual(len(self.repository.find_by_hash('org/any', hashlib.sha256(file.read_bytes()).hexdigest())), 1)

    def test_external_key_resolves_new_version_but_filename_and_title_do_not(self):
        first_file = self.docx('same.docx', ('Rule version one.',))
        a = self.ingest(first_file, external_key='external/opaque/key')
        second_file = self.docx('same.docx', ('Rule version two.',))
        b = self.ingest(second_file, external_key='external/opaque/key')
        self.assertEqual(a.document.identity, b.document.identity)
        self.assertNotEqual(a.version.identity, b.version.identity)
        c = self.ingest(self.docx('same.docx', ('Unrelated third source.',)))
        self.assertNotEqual(a.document.identity, c.document.identity)

    def test_existing_arbitrary_document_id_resolves_a_new_draft(self):
        a = self.ingest(document_id='opaque/id with spaces Ω')
        b = self.ingest(self.docx('other.docx', ('Completely different content.',)), document_id=a.document.identity.document_id)
        self.assertEqual(a.document.identity, b.document.identity)
        self.assertNotEqual(a.version.identity, b.version.identity)
        self.assertEqual(b.version.status, DocumentVersionStatus.DRAFT)

    def test_conflicting_external_key_and_document_id_are_rejected(self):
        self.ingest(external_key='key', document_id='d1')
        with self.assertRaises(IngestionError):
            self.ingest(self.docx('other.docx', ('Different content.',)), external_key='key', document_id='d2')
        self.assertEqual(self.table_count('documents'), 1)

    def test_cross_organization_id_cannot_alias_or_expose_duplicate(self):
        file = self.docx()
        a = self.ingest(file, document_id='global/document')
        with self.assertRaises(IngestionError):
            self.service.ingest_new_document(file, title='Different tenant', organization_id='foreign',
                                             version_label='v', document_id=a.document.identity.document_id)
        self.assertEqual(self.table_count('versions'), 1)

    def test_old_active_survives_new_draft_then_is_superseded_atomically(self):
        a = self.ingest()
        self.publish(a)
        b = self.service.ingest_new_version(self.docx('updated.docx', ('New independent rule.',)),
                                           document=a.document.identity, version_label='revision/second')
        self.assertEqual(b.version.status, DocumentVersionStatus.DRAFT)
        self.assertEqual(self.repository.get_current_version(a.document.identity).identity, a.version.identity)
        self.publish(b, a.version.identity.version_id)
        self.assertEqual(self.repository.get_version(a.version.identity).status, DocumentVersionStatus.SUPERSEDED)
        self.assertEqual(self.repository.get_current_version(a.document.identity).identity, b.version.identity)
        self.assertEqual(self.repository.get_version(b.version.identity).processing_state, ProcessingState.CHUNKED)
        self.assertTrue(self.repository.get_version(b.version.identity).approved)
        self.assertTrue(self.repository.list_snapshot(a.snapshot))
        self.assertEqual(self.table_count('versions'), 2)

    def test_archive_excludes_current_without_deleting_history(self):
        result = self.ingest()
        self.publish(result)
        self.service.archive_version(result.version.identity)
        self.assertIsNone(self.repository.get_current_version(result.document.identity))
        self.assertEqual(self.repository.get_version(result.version.identity).status, DocumentVersionStatus.ARCHIVED)
        self.assertTrue(self.repository.list_snapshot(result.snapshot))
        with self.assertRaises(IngestionError):
            self.publish(result)

    def test_activation_checks_expected_current_across_connections(self):
        a = self.ingest()
        b = self.service.ingest_new_version(self.docx('next.docx', ('New contents.',)),
                                           document=a.document.identity, version_label='next')
        with SQLiteRepository(self.database) as other:
            other.activate_version(a.version.identity, expected_current_version_id=None)
        with self.assertRaises(IngestionError):
            self.publish(b, None)
        self.assertEqual(self.repository.get_current_version(a.document.identity).identity, a.version.identity)
        self.assertEqual(self.repository.get_version(b.version.identity).status, DocumentVersionStatus.DRAFT)

    def test_activation_database_failure_rolls_back_superseding(self):
        a = self.ingest()
        self.publish(a)
        b = self.service.ingest_new_version(self.docx('next.docx', ('New contents.',)),
                                           document=a.document.identity, version_label='next')
        self.repository.connection.execute('''CREATE TEMP TRIGGER reject_activation BEFORE UPDATE OF status ON versions
            WHEN NEW.status='active' BEGIN SELECT RAISE(ABORT,'PRIVATE_DATABASE_DETAIL'); END''')
        with self.assertRaises(IngestionError) as error:
            self.publish(b, a.version.identity.version_id)
        self.assertNotIn('PRIVATE', str(error.exception))
        self.assertEqual(self.repository.get_current_version(a.document.identity).identity, a.version.identity)
        self.assertEqual(self.repository.get_version(b.version.identity).status, DocumentVersionStatus.DRAFT)

    def test_published_chunked_versions_are_not_yet_retrieval_visible(self):
        result = self.ingest(required_scopes=frozenset({'read'}))
        self.publish(result)
        access = AccessContext('org/any', 'principal', {'read'})
        self.assertEqual(self.repository.list_visible_versions(access, filters=RetrievalFilters()), ())

    def test_repository_visibility_filters_organization_scopes_and_lifecycle(self):
        a = self.ingest(required_scopes=frozenset({'read'}))
        self.publish(a)
        current = self.repository.get_version(a.version.identity)
        # Simulate a future indexing completion flag; no vector/model execution.
        self.repository.save_version(replace(current, processing_state=ProcessingState.INDEXED))
        access = AccessContext('org/any', 'reader', {'read'})
        self.assertEqual(len(self.repository.list_visible_versions(access, filters=RetrievalFilters())), 1)
        self.assertEqual(self.repository.list_visible_versions(AccessContext('foreign', 'reader', {'read'}), filters=RetrievalFilters()), ())
        self.assertEqual(self.repository.list_visible_versions(AccessContext('org/any', 'reader'), filters=RetrievalFilters()), ())
        self.assertEqual(self.repository.list_visible_versions(access, filters=RetrievalFilters(allowed_document_ids=frozenset())), ())
        self.assertEqual(self.repository.list_visible_versions(access, filters=RetrievalFilters(allowed_version_ids={'other'})), ())
        self.service.archive_version(a.version.identity)
        self.assertEqual(self.repository.list_visible_versions(access, filters=RetrievalFilters()), ())

    def test_preflight_empty_missing_and_unsupported_files_create_no_records(self):
        cases = [(self.path / 'missing.docx', 'file_not_found')]
        empty = self.path / 'empty.docx'
        empty.write_bytes(b'')
        unsupported = self.path / 'input.txt'
        unsupported.write_text('arbitrary text')
        cases.extend([(empty, 'empty_file'), (unsupported, 'unsupported_format')])
        for file, reason in cases:
            with self.subTest(reason=reason), self.assertRaises(IngestionError) as error:
                self.ingest(file)
            self.assertEqual(error.exception.reason_code, reason)
            self.assertIsNone(error.exception.version)
        self.assertEqual(self.table_count('documents'), 0)
        self.assertEqual(self.table_count('versions'), 0)

    def test_parsing_failure_records_failed_and_preserves_previous_active(self):
        a = self.ingest()
        self.publish(a)
        broken = self.path / 'broken.docx'
        broken.write_bytes(b'not a document PRIVATE_SOURCE')
        with self.assertRaises(IngestionError) as error:
            self.service.ingest_new_version(broken, document=a.document.identity, version_label='broken')
        failure = error.exception
        self.assertEqual(failure.reason_code, 'parsing_failed')
        self.assertTrue(failure.failure_recorded)
        failed = self.repository.get_version(failure.version)
        self.assertEqual(failed.processing_state, ProcessingState.FAILED)
        self.assertEqual(failed.status, DocumentVersionStatus.DRAFT)
        self.assertEqual(failed.metadata['failure_reason'], 'parsing_failed')
        self.assertNotIn('PRIVATE', str(failure) + str(failed.metadata))
        self.assertEqual(self.repository.get_current_version(a.document.identity).identity, a.version.identity)
        with self.assertRaises(IngestionError):
            self.service.activate_version(failed.identity, expected_current_version_id=a.version.identity.version_id)

    def test_empty_extraction_is_failed_not_prepared(self):
        with self.assertRaises(IngestionError) as error:
            self.ingest(self.docx('blank.docx', ()))
        self.assertEqual(error.exception.reason_code, 'empty_extraction')
        self.assertEqual(self.repository.get_version(error.exception.version).processing_state, ProcessingState.FAILED)
        self.assertEqual(self.table_count('snapshots'), 0)

    def test_chunking_failure_keeps_old_active_and_sanitizes_diagnostics(self):
        a = self.ingest()
        self.publish(a)
        with patch('knowledge_base.chunker.chunk_source_blocks', side_effect=RuntimeError('SECRET_TOKEN')):
            with self.assertRaises(IngestionError) as error:
                self.service.ingest_new_version(self.docx('new.docx', ('New version contents.',)),
                                               document=a.document.identity, version_label='new')
        failure = error.exception
        self.assertEqual(failure.reason_code, 'chunking_failed')
        self.assertTrue(failure.failure_recorded)
        self.assertNotIn('SECRET', str(failure) + str(self.repository.get_version(failure.version).metadata))
        self.assertEqual(self.repository.get_current_version(a.document.identity).identity, a.version.identity)
        self.assertEqual(self.table_count('snapshots'), 1)

    def test_empty_chunker_output_is_failed(self):
        with patch('knowledge_base.chunker.chunk_source_blocks', return_value=()):
            with self.assertRaises(IngestionError) as error:
                self.ingest()
        self.assertEqual(error.exception.reason_code, 'empty_chunks')
        self.assertTrue(error.exception.failure_recorded)
        self.assertEqual(self.table_count('chunks'), 0)

    def test_chunk_write_failure_rolls_back_snapshot_and_partial_chunks(self):
        a = self.ingest()
        self.publish(a)
        previous_chunks = self.table_count('chunks')
        self.repository.connection.execute('''CREATE TEMP TRIGGER reject_second_chunk BEFORE INSERT ON chunks
            WHEN NEW.ordinal=1 BEGIN SELECT RAISE(ABORT,'PRIVATE_INSERT_DETAIL'); END''')
        file = self.docx('long.docx', ('A separate rule. ' * 50,))
        with self.assertRaises(IngestionError) as error:
            self.service.ingest_new_version(file, document=a.document.identity, version_label='long')
        failure = error.exception
        self.assertEqual(failure.reason_code, 'storage_failed')
        self.assertTrue(failure.failure_recorded)
        self.assertIsNone(self.repository.get_current_snapshot(failure.version))
        self.assertEqual(self.table_count('chunks'), previous_chunks)
        self.assertEqual(self.table_count('snapshots'), 1)
        self.assertEqual(self.repository.get_current_version(a.document.identity).identity, a.version.identity)

    def test_registration_failure_leaves_no_orphan_document_or_version(self):
        self.repository.connection.execute('''CREATE TEMP TRIGGER reject_version BEFORE INSERT ON versions
            BEGIN SELECT RAISE(ABORT,'PRIVATE_INSERT_DETAIL'); END''')
        with self.assertRaises(IngestionError) as error:
            self.ingest()
        self.assertEqual(error.exception.reason_code, 'storage_failed')
        self.assertIsNone(error.exception.version)
        self.assertEqual(self.table_count('documents'), 0)
        self.assertEqual(self.table_count('versions'), 0)

    def test_unavailable_database_reports_failure_without_claiming_persistence(self):
        with patch.object(self.repository, 'finish_processing', side_effect=StorageError('PRIVATE_DETAIL')), \
             patch.object(self.repository, 'fail_processing', side_effect=StorageError('PRIVATE_DETAIL')):
            with self.assertRaises(IngestionError) as error:
                self.ingest()
        self.assertFalse(error.exception.failure_recorded)
        self.assertEqual(str(error.exception), 'storage_failed')
        self.assertEqual(self.repository.get_version(error.exception.version).processing_state, ProcessingState.PARSED)
        self.assertEqual(self.table_count('snapshots'), 0)

    def test_restart_preserves_source_identity_chunks_states_and_nested_metadata(self):
        result = self.ingest(metadata={'nested': {'values': [1, 'Ω', {'flag': True}]}})
        self.publish(result)
        chunks = self.repository.list_snapshot(result.snapshot)
        source = self.repository.get_source(result.version.identity)
        self.repository.close()
        with SQLiteRepository(self.database) as reopened:
            self.assertEqual(reopened.get_document(result.document.identity), result.document)
            self.assertEqual(reopened.list_snapshot(result.snapshot), chunks)
            self.assertEqual(reopened.get_current_version(result.document.identity).identity, result.version.identity)
            self.assertEqual(reopened.get_source(result.version.identity), source)
            # Recreating a DOCX can change ZIP timestamps; copy the saved source explicitly.
            original = self.path / 'original.docx'
            original.write_bytes(reopened.get_source(result.version.identity))
            duplicate = IngestionService(reopened).ingest_new_version(original, document=result.document.identity, version_label='duplicate')
            self.assertTrue(duplicate.duplicate)
            self.assertEqual(duplicate.snapshot, result.snapshot)

    def test_snapshot_is_idempotent_immutable_and_ids_are_checked(self):
        result = self.ingest()
        chunks = self.repository.list_snapshot(result.snapshot)
        self.repository.save_snapshot(result.snapshot, chunks)
        self.assertEqual(self.repository.list_snapshot(result.snapshot), chunks)
        with self.assertRaises(StorageError):
            self.repository.save_snapshot(result.snapshot, [replace(chunks[0], text='changed')])
        wrong = replace(result.snapshot, processing_snapshot_id='other')
        with self.assertRaises(StorageError):
            self.repository.save_snapshot(wrong, chunks)
        with self.assertRaises(StorageError):
            self.repository.delete_snapshot(result.snapshot)
        self.assertIsNone(self.repository.get_chunk(replace(chunks[0].identity, version_id='foreign')))

    def test_direct_lifecycle_changes_and_processing_shortcuts_are_rejected(self):
        document = RuntimeDocument(DocumentIdentity('arbitrary/direct', 'org/direct'), 'Neutral title')
        version = RuntimeDocumentVersion(DocumentVersionIdentity(document.identity.document_id, 'version#custom'), 'custom')
        self.repository.save_document(document)
        self.repository.save_version(version)
        with self.assertRaises(StorageError):
            self.repository.save_version(replace(version, status=DocumentVersionStatus.ACTIVE))
        with self.assertRaises(StorageError):
            self.repository.save_version(replace(version, processing_state=ProcessingState.CHUNKED))
        with self.assertRaises(StorageError):
            self.repository.activate_version(version.identity, expected_current_version_id=None)

    def test_source_bytes_are_immutable_for_an_existing_version_identity(self):
        result = self.ingest()
        version = self.repository.get_version(result.version.identity)
        with self.assertRaises(StorageError):
            self.repository.register_source(result.document, replace(version, processing_state=ProcessingState.REGISTERED),
                source=b'changed bytes', file_name='changed.docx', source_format='docx', external_key=None, explicit_identity=True)
        self.assertEqual(self.repository.get_source_info(version.identity).sha256,
                         hashlib.sha256(self.repository.get_source(version.identity)).hexdigest())

    def test_source_changed_between_registration_and_read_is_rejected(self):
        file = self.docx()
        from knowledge_base.document_registry import register_document
        registered = register_document(file)
        file.write_bytes(b'changed')
        with patch('knowledge_base.runtime.ingestion.register_document', return_value=registered):
            with self.assertRaises(IngestionError) as error:
                self.ingest(file)
        self.assertEqual(error.exception.reason_code, 'source_changed_during_read')
        self.assertEqual(self.table_count('versions'), 0)

    def test_runtime_imports_never_reference_benchmark_or_model_execution(self):
        allowed = {'knowledge_base.document_intake', 'knowledge_base.document_registry',
                   'knowledge_base.document_versions', 'knowledge_base.source_blocks', 'knowledge_base.chunker'}
        for path in (ROOT / 'src/knowledge_base/runtime').glob('*.py'):
            for node in ast.walk(ast.parse(path.read_text(encoding='utf8'))):
                if isinstance(node, ast.ImportFrom) and node.level == 0 and node.module.startswith('knowledge_base.'):
                    self.assertIn(node.module, allowed)
        self.assertNotIn('expected_behavior', inspect.signature(IngestionService.ingest_new_document).parameters)
        self.assertNotIn('query_id', inspect.signature(IngestionService.ingest_new_version).parameters)

    def test_processing_state_write_failure_rolls_back_all_new_chunks(self):
        old = self.ingest()
        self.publish(old)
        count = self.table_count('chunks')
        self.repository.connection.execute('''CREATE TEMP TRIGGER reject_chunked BEFORE UPDATE OF processing_state ON versions
            WHEN NEW.processing_state='chunked' BEGIN SELECT RAISE(ABORT,'PRIVATE_STATE_DETAIL'); END''')
        with self.assertRaises(IngestionError) as error:
            self.service.ingest_new_version(self.docx('updated.docx', ('New prepared contents.',)),
                                           document=old.document.identity, version_label='next')
        self.assertTrue(error.exception.failure_recorded)
        self.assertEqual(self.repository.get_version(error.exception.version).processing_state, ProcessingState.FAILED)
        self.assertIsNone(self.repository.get_current_snapshot(error.exception.version))
        self.assertEqual(self.table_count('chunks'), count)
        self.assertEqual(self.table_count('snapshots'), 1)
        self.assertEqual(self.repository.get_current_version(old.document.identity).identity, old.version.identity)

    def test_repeated_failed_source_is_duplicate_not_silently_reprocessed(self):
        file = self.docx('empty.docx', ())
        with self.assertRaises(IngestionError) as error:
            self.ingest(file)
        again = self.ingest(file)
        self.assertTrue(again.duplicate)
        self.assertEqual(again.version.identity, error.exception.version)
        self.assertEqual(again.version.processing_state, ProcessingState.FAILED)
        self.assertEqual(self.table_count('versions'), 1)

    def test_saved_snapshot_that_is_not_selected_can_be_deleted(self):
        result = self.ingest()
        chunk = self.repository.list_snapshot(result.snapshot)[0]
        snapshot = replace(result.snapshot, processing_snapshot_id='unselected/opaque')
        other = replace(chunk, identity=replace(chunk.identity,
            processing_snapshot_id=snapshot.processing_snapshot_id, chunk_id='different/opaque'))
        self.repository.save_snapshot(snapshot, [other])
        self.assertEqual(self.repository.get_chunk(other.identity), other)
        self.repository.delete_snapshot(snapshot)
        self.assertIsNone(self.repository.get_chunk(other.identity))
        self.assertTrue(self.repository.list_snapshot(result.snapshot))

    def test_external_key_race_does_not_override_an_explicit_document_id(self):
        existing = self.ingest(external_key='race-key', document_id='existing/id')
        pending = RuntimeDocument(DocumentIdentity('requested/id', 'org/any'), 'Other logical document')
        version = RuntimeDocumentVersion(DocumentVersionIdentity('requested/id', 'new/version'), 'new')
        with self.assertRaises(StorageError):
            self.repository.register_source(pending, version, source=b'new source', file_name='new.docx',
                source_format='docx', external_key='race-key', explicit_identity=True)
        self.assertEqual(self.table_count('documents'), 1)
        self.assertEqual(self.table_count('versions'), 1)

    def test_external_key_race_can_resolve_a_provisional_generated_identity(self):
        first = self.ingest(external_key='race-key')
        pending = RuntimeDocument(DocumentIdentity('provisional/id', 'org/any'), 'Provisional document')
        version = RuntimeDocumentVersion(DocumentVersionIdentity('provisional/id', 'new/version'), 'new')
        actual, duplicate = self.repository.register_source(pending, version, source=b'new source', file_name='new.docx',
            source_format='docx', external_key='race-key', explicit_identity=True, resolve_external_key=True)
        self.assertFalse(duplicate)
        self.assertEqual(actual.identity.document_id, first.document.identity.document_id)
        self.assertEqual(self.table_count('documents'), 1)
        self.assertEqual(self.table_count('versions'), 2)


if __name__ == '__main__':
    unittest.main()
