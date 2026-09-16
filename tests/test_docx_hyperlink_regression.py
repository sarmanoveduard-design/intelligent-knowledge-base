import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from knowledge_base.docx_normalizer import normalize_docx


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        RT.HYPERLINK,
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(
        qn("r:id"),
        relationship_id,
    )

    run = OxmlElement("w:r")

    run_properties = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(
        qn("w:val"),
        "0563C1",
    )

    underline = OxmlElement("w:u")
    underline.set(
        qn("w:val"),
        "single",
    )

    run_properties.append(color)
    run_properties.append(underline)

    run.append(run_properties)

    text_element = OxmlElement("w:t")
    text_element.text = text

    run.append(text_element)
    hyperlink.append(run)

    paragraph._p.append(hyperlink)


class DocxHyperlinkRegressionTests(unittest.TestCase):
    def test_hyperlink_text_is_not_lost(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "law.docx"

            document = Document()

            paragraph = document.add_paragraph()

            paragraph.add_run(
                "Уголовный кодекс Российской Федерации "
            )

            add_hyperlink(
                paragraph,
                "от 13.06.1996 N 63-ФЗ",
                "https://example.com/law",
            )

            paragraph.add_run(
                " (ред. от 21.11.2022)"
            )

            document.save(file_path)

            lines = normalize_docx(file_path)

            self.assertEqual(len(lines), 1)

            self.assertEqual(
                lines[0].text,
                (
                    "Уголовный кодекс Российской Федерации "
                    "от 13.06.1996 N 63-ФЗ "
                    "(ред. от 21.11.2022)"
                ),
            )

            self.assertTrue(
                lines[0].formatting_reliable
            )


if __name__ == "__main__":
    unittest.main()
