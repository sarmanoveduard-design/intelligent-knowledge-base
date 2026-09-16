import tempfile
import unittest
from pathlib import Path

from docx import Document

from knowledge_base.source_blocks import (
    SourceBlock,
    build_source_blocks,
)


class SourceBlockTests(unittest.TestCase):
    def test_source_block_keeps_document_identity(self):
        block = SourceBlock(
            document_id="doc-1",
            version_id="ver-1",
            block_index=0,
            text="Тест",
            source_file="test.docx",
            source_format="docx",
        )

        self.assertEqual(
            block.document_id,
            "doc-1",
        )

        self.assertEqual(
            block.version_id,
            "ver-1",
        )

    def test_docx_becomes_source_blocks(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "sample.docx"

            document = Document()
            document.add_paragraph("Первый абзац")
            document.add_paragraph("Второй абзац")
            document.save(file_path)

            blocks = build_source_blocks(
                file_path,
                document_id="doc-1",
                version_id="ver-1",
            )

            self.assertEqual(len(blocks), 2)

            self.assertEqual(
                blocks[0].text,
                "Первый абзац",
            )

            self.assertEqual(
                blocks[1].text,
                "Второй абзац",
            )

            self.assertEqual(
                blocks[0].source_format,
                "docx",
            )

            self.assertEqual(
                blocks[0].paragraph_index,
                0,
            )

            self.assertIsNone(
                blocks[0].page_number,
            )

    def test_docx_manual_lines_keep_source_location(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "lines.docx"

            document = Document()

            paragraph = document.add_paragraph()
            paragraph.add_run("Строка один")
            paragraph.add_run().add_break()
            paragraph.add_run("Строка два")

            document.save(file_path)

            blocks = build_source_blocks(
                file_path,
                document_id="doc-1",
                version_id="ver-1",
            )

            self.assertEqual(len(blocks), 2)

            self.assertEqual(
                blocks[0].paragraph_index,
                blocks[1].paragraph_index,
            )

            self.assertEqual(
                blocks[0].line_index,
                0,
            )

            self.assertEqual(
                blocks[1].line_index,
                1,
            )

    def test_unsupported_file_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "sample.txt"
            file_path.write_text(
                "test",
                encoding="utf-8",
            )

            with self.assertRaises(ValueError):
                build_source_blocks(
                    file_path,
                    document_id="doc-1",
                    version_id="ver-1",
                )


if __name__ == "__main__":
    unittest.main()
