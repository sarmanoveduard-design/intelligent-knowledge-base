import tempfile
import unittest
from pathlib import Path

from docx import Document

from knowledge_base.docx_parser import parse_docx


class DocxParserTests(unittest.TestCase):
    def test_extracts_text_from_real_docx(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "lecture.docx"

            document = Document()
            document.add_heading(
                "Управление корпоративным здоровьем",
                level=1,
            )
            document.add_paragraph(
                "Первый содержательный абзац."
            )
            document.add_paragraph(
                "Второй содержательный абзац."
            )
            document.save(file_path)

            parsed = parse_docx(file_path)

            self.assertEqual(
                parsed.file_name,
                "lecture.docx",
            )

            self.assertEqual(
                len(parsed.blocks),
                3,
            )

            self.assertEqual(
                parsed.blocks[0].block_type,
                "heading",
            )

            self.assertEqual(
                parsed.blocks[0].text,
                "Управление корпоративным здоровьем",
            )

            self.assertIn(
                "Первый содержательный абзац.",
                parsed.full_text,
            )

    def test_empty_paragraphs_are_ignored(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "empty.docx"

            document = Document()
            document.add_paragraph("")
            document.add_paragraph("Полезный текст")
            document.add_paragraph("   ")
            document.save(file_path)

            parsed = parse_docx(file_path)

            self.assertEqual(
                len(parsed.blocks),
                1,
            )

            self.assertEqual(
                parsed.blocks[0].text,
                "Полезный текст",
            )

    def test_rejects_broken_docx(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "broken.docx"
            file_path.write_bytes(
                b"this is not a real Word document"
            )

            with self.assertRaises(ValueError):
                parse_docx(file_path)


if __name__ == "__main__":
    unittest.main()
