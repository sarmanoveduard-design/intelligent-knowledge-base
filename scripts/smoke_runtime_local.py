"""Small local synthetic smoke, not a quality benchmark or model downloader.

Run only after offline unit suites. Uses the existing optional reranker runtime
and cached weights; no dependencies, model downloads, or cloud calls are made.
"""
import argparse
from dataclasses import asdict
from datetime import datetime, timezone
import json
import os
from pathlib import Path
from time import perf_counter

from knowledge_base.ollama_embeddings import OllamaEmbeddingProvider
from knowledge_base.runtime.composition import compose_local_runtime
from knowledge_base.runtime.config import RuntimeConfig
from knowledge_base.runtime.models import AccessContext, AnswerState, AskRequest
from knowledge_base.runtime.ollama import OllamaGenerationConfig, check_local_models
from knowledge_base.runtime.retrieval import BGERuntimeAdapter


def _cached_reranker(cache_dir, device):
    # Existing implementation remains untouched. Both loaders are explicitly
    # local-only, even if a caller's HF_HUB_OFFLINE was previously disabled.
    os.environ['HF_HUB_OFFLINE'] = '1'
    os.environ['TRANSFORMERS_OFFLINE'] = '1'
    os.environ['HF_HUB_DISABLE_TELEMETRY'] = '1'
    os.environ['HF_HUB_DISABLE_IMPLICIT_TOKEN'] = '1'
    try:
        import torch
        from transformers import AutoModelForSequenceClassification, AutoTokenizer
        from knowledge_base.bge_reranker import BGEReranker, BGERerankerConfig, _TransformersBackend
        class LocalOnly:
            def __init__(self, loader): self.loader = loader
            def from_pretrained(self, *args, **kwargs):
                return self.loader.from_pretrained(*args, **kwargs, local_files_only=True)
        config = BGERerankerConfig(device=device, cache_dir=str(cache_dir))
        backend = _TransformersBackend(config, torch, LocalOnly(AutoTokenizer), LocalOnly(AutoModelForSequenceClassification))
        legacy = BGEReranker(config, backend=backend)
        return BGERuntimeAdapter(legacy), legacy.metadata
    except Exception:
        raise RuntimeError('LOCAL_RERANKER_UNAVAILABLE') from None


def run_smoke(output_dir, *, base_url, model, embedding_model, reranker_cache, reranker_device):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=False)
    config = OllamaGenerationConfig(model=model, base_url=base_url, think=False, temperature=0.0,
                                    seed=0, timeout_seconds=360, num_predict=512, num_ctx=8192, keep_alive='30m')
    report = {'started_at': datetime.now(timezone.utc).isoformat(), 'generation_config': asdict(config),
              'embedding_model': embedding_model, 'questions': [], 'no_model_downloads': True,
              'warmup_used': False}
    began = perf_counter()
    phase = 'OLLAMA_PREFLIGHT_FAILED'
    try:
        inventory = check_local_models(config, embedding_model)
        report['ollama'] = asdict(inventory)
        report['ollama']['ready'] = inventory.ready
        if not inventory.ready:
            phase = 'LOCAL_MODEL_MISSING'
            raise RuntimeError(phase)
        phase = 'LOCAL_RERANKER_UNAVAILABLE'
        started = perf_counter()
        reranker, metadata = _cached_reranker(reranker_cache, reranker_device)
        report['reranker'] = metadata
        report['reranker_load_seconds'] = perf_counter() - started
        from docx import Document
        document = Document()
        document.add_heading('Правила офиса Альфа', level=1)
        for text in ('Офис открывается в 08:30.', 'Гостевой пропуск действует 2 часа.',
                     'Серверная доступна только сотрудникам группы IT.'):
            document.add_paragraph(text)
        source = output_dir / 'office-alpha.docx'
        document.save(source)
        phase = 'RUNTIME_SMOKE_FAILED'
        embeddings = OllamaEmbeddingProvider(base_url=base_url, model=embedding_model, dimension=1024, timeout=180)
        with compose_local_runtime(output_dir / 'registry.sqlite', embeddings=embeddings, reranker=reranker,
            generation_config=config, runtime_config=RuntimeConfig(candidate_top_k=3, final_top_k=2)) as runtime:
            started = perf_counter()
            ingested = runtime.ingestion.ingest_new_document(source, title='Правила офиса Альфа',
                organization_id='synthetic/alpha', version_label='v1', required_scopes={'office.read'})
            report['ingestion_seconds'] = perf_counter() - started
            started = perf_counter()
            manifest = runtime.indexing.index_version(ingested.version.identity)
            runtime.ingestion.activate_version(ingested.version.identity, expected_current_version_id=None)
            report['indexing_seconds'] = perf_counter() - started
            report['source_file'] = source.name
            report['chunk_count'] = manifest.chunk_count
            cases = (
                ('answerable', 'Во сколько открывается офис?', {'office.read'}),
                ('missing_fact', 'Какой пароль от Wi-Fi?', {'office.read'}),
                ('no_scope', 'Кому разрешён доступ к серверной?', set()),
            )
            observed = {}
            original_generate = runtime.generation.generate
            def observe_generation(request):
                draft = original_generate(request)
                observed['draft'] = {'answer': draft.answer, 'declared_limitations': list(draft.declared_limitations),
                                     'citations': [asdict(x) for x in draft.citations]}
                return draft
            runtime.generation.generate = observe_generation
            for position, (label, question, scopes) in enumerate(cases, start=1):
                observed.clear()
                request = AskRequest('local-smoke/' + str(position), question,
                    AccessContext('synthetic/alpha', 'synthetic/reader', scopes), language='ru')
                started = perf_counter()
                result = runtime.ask(request)
                called = result.diagnostics.get('generation', {}).get('called', False)
                validation = result.citation_validation
                report['questions'].append({
                    'case': label, 'question': question, 'request_id': result.request_id,
                    'state': result.state.value, 'reason_code': result.reason_code, 'answer': result.answer,
                    'declared_limitations': list(result.declared_limitations),
                    'source_handles': [x.source_handle for x in result.context_manifest.entries] if result.context_manifest else [],
                    'citations': [x.source_handle for x in validation.resolved_citations] if validation else [],
                    'structural_validity': validation.structural_validity.value if validation else
                        ('not_applicable' if result.diagnostics.get('citation', {}).get('status') == 'not_applicable' else 'not_run'),
                    'semantic_support': validation.semantic_support.value if validation else 'not_checked',
                    'citation_errors': list(validation.errors) if validation else [],
                    'generation_called': called, 'elapsed_seconds': perf_counter() - started,
                    'generation_draft': observed.get('draft'),
                    'generation_metadata': dict(runtime.generation.last_call_metadata) if called else {},
                })
            # Explicit restart demonstrates that no memory vectors survive.
        with compose_local_runtime(output_dir / 'registry.sqlite', embeddings=embeddings, reranker=reranker,
            generation_config=config) as reopened:
            access = AccessContext('synthetic/alpha', 'synthetic/reader', {'office.read'})
            report['readiness_after_restart'] = asdict(reopened.readiness(access))
            started = perf_counter()
            rebuilt = reopened.rebuild(access)
            report['rebuild_seconds'] = perf_counter() - started
            report['rebuilt_snapshots'] = len(rebuilt)
            report['readiness_after_rebuild'] = asdict(reopened.readiness(access))
        report['status'] = 'completed'
        questions = report['questions']
        report['technical_smoke_passed'] = (
            questions[0]['state'] == AnswerState.ANSWER.value and questions[0]['structural_validity'] == 'pass'
            and '08:30' in (questions[0]['answer'] or '') and questions[0]['citations'] == ['S1']
            and questions[0]['generation_called']
            and questions[1]['state'] == AnswerState.REFUSE_INSUFFICIENT_CONTEXT.value
            and questions[1]['reason_code'] == 'GENERATION_REPORTED_INSUFFICIENT_CONTEXT'
            and questions[1]['generation_called'] and questions[1]['generation_draft'] is not None
            and questions[1]['generation_draft']['answer'] == ''
            and bool(questions[1]['generation_draft']['declared_limitations'])
            and questions[1]['generation_draft']['citations'] == []
            and questions[1]['citations'] == [] and questions[1]['citation_errors'] == []
            and questions[1]['structural_validity'] == 'not_applicable'
            and questions[2]['state'] == AnswerState.REFUSE_INSUFFICIENT_CONTEXT.value
            and questions[2]['reason_code'] == 'NO_EVIDENCE'
            and not questions[2]['generation_called'])
    except Exception:
        report['status'] = 'stopped'
        report['reason_code'] = phase
    report['total_elapsed_seconds'] = perf_counter() - began
    (output_dir / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf8')
    print(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False))
    return 0 if report.get('technical_smoke_passed') else 2


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', required=True)
    parser.add_argument('--base-url', default='http://localhost:11434')
    parser.add_argument('--model', default='qwen3:8b')
    parser.add_argument('--embedding-model', default='bge-m3')
    parser.add_argument('--reranker-cache', required=True)
    parser.add_argument('--reranker-device', choices=('cpu', 'cuda', 'auto'), default='cpu')
    args = parser.parse_args()
    return run_smoke(args.output_dir, base_url=args.base_url, model=args.model,
        embedding_model=args.embedding_model, reranker_cache=args.reranker_cache, reranker_device=args.reranker_device)


if __name__ == '__main__':
    raise SystemExit(main())
