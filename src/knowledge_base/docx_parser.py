from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document


@dataclass(frozen=True)
class TextBlock:
    block_index: int
    text: str
    style: str | None
    block_type: str


@dataclass(frozen=True)
class ParsedDocument:
    file_name: str
    blocks: tuple[TextBlock, ...]

    @property
    def full_text(self) -> str:
        return "\n\n".join(
            block.text
            for block in self.blocks
            if block.text
        )


def classify_paragraph(style_name: str | None) -> str:
    if style_name and style_name.lower().startswith("heading"):
        return "heading"

    return "paragraph"


def parse_docx(file_path: Path) -> ParsedDocument:
    if not file_path.is_file():
        raise FileNotFoundError(
            f"Файл не найден: {file_path}"
        )

    if file_path.suffix.lower() != ".docx":
        raise ValueError(
            "DOCX-парсер принимает только .docx файлы"
        )

    try:
        document = Document(file_path)
    except Exception as exc:
        raise ValueError(
            f"Не удалось прочитать DOCX: {file_path.name}"
        ) from exc

    blocks: list[TextBlock] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        style_name = (
            paragraph.style.name
            if paragraph.style is not None
            else None
        )

        blocks.append(
            TextBlock(
                block_index=len(blocks),
                text=text,
                style=style_name,
                block_type=classify_paragraph(style_name),
            )
        )

    return ParsedDocument(
        file_name=file_path.name,
        blocks=tuple(blocks),
    )
